# -*- coding: utf-8 -*-
"""Build the 14-slide judge deck: navy/yellow, large type, one idea per slide."""
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
SHOTS = HERE / "shots"
DOCS = Path(r"C:\Projects\Trading\docs\screenshots")

WHITE = RGBColor(0xFA, 0xFA, 0xFA)
YELLOW = RGBColor(0xFF, 0xC6, 0x1A)
MUTED = RGBColor(0xA8, 0xBC, 0xE0)
CARD = RGBColor(0x0B, 0x1F, 0x4E)
CARD_LINE = RGBColor(0x27, 0x44, 0x86)
SANS = "Segoe UI"
MONO = "Consolas"
URL = "alpaca-trade-intelli.vercel.app"

prs = Presentation()
prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
BLANK = prs.slide_layouts[6]

W, H = 13.3333, 7.5
M = 0.78
CW = W - 2 * M


def slide():
    s = prs.slides.add_slide(BLANK)
    s.shapes.add_picture(str(HERE / "bg.jpg"), 0, 0, prs.slide_width, prs.slide_height)
    return s


def text(s, x, y, w, h, runs, size, bold=False, color=WHITE, font=SANS,
         align=PP_ALIGN.LEFT, spacing=1.0, anchor=MSO_ANCHOR.TOP):
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = spacing
    for t, c, b in ([(runs, color, bold)] if isinstance(runs, str) else runs):
        r = p.add_run()
        r.text = t
        r.font.size = Pt(size)
        r.font.bold = b
        r.font.color.rgb = c
        r.font.name = font
    return box


def card(s, x, y, w, h, fill=CARD, line=CARD_LINE, width=1.0):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                            Inches(x), Inches(y), Inches(w), Inches(h))
    sh.adjustments[0] = 0.05
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = line
    sh.line.width = Pt(width)
    sh.shadow.inherit = False
    sh.text_frame.text = ""
    return sh


def eyebrow(s, label):
    text(s, M, 0.52, CW, 0.3, label.upper(), 13, bold=True, color=YELLOW)


def heading(s, white, yellow=""):
    runs = [(white, WHITE, True)] + ([(yellow, YELLOW, True)] if yellow else [])
    text(s, M, 0.9, CW, 0.9, runs, 40, spacing=1.05)


def footer(s, line, y=6.45):
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(M), Inches(y),
                             Inches(0.055), Inches(0.44))
    bar.fill.solid()
    bar.fill.fore_color.rgb = YELLOW
    bar.line.fill.background()
    bar.shadow.inherit = False
    text(s, M + 0.26, y, CW - 0.26, 0.44, line, 19, color=YELLOW, anchor=MSO_ANCHOR.MIDDLE)


def rule(s, y, w=CW, x=M, color=YELLOW, h=0.022):
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    sh.line.fill.background()
    sh.shadow.inherit = False


def columns(s, items, y, h, render, gap=0.35):
    n = len(items)
    cw = (CW - gap * (n - 1)) / n
    for i, item in enumerate(items):
        x = M + i * (cw + gap)
        card(s, x, y, cw, h)
        render(s, x, y, cw, item)


def picture(s, name, y, caption=None, max_bottom=5.86):
    """Widest the image can be without colliding with the caption and footer."""
    iw, ih = Image.open(SHOTS / name).size
    width = min(CW, (max_bottom - y) * iw / ih)
    x = (W - width) / 2
    pic = s.shapes.add_picture(str(SHOTS / name), Inches(x), Inches(y), width=Inches(width))
    pic.line.color.rgb = YELLOW
    pic.line.width = Pt(1.0)
    if caption:
        text(s, x, Emu(pic.top + pic.height).inches + 0.2, width, 0.4, caption, 17, color=MUTED)
    return pic


def crop(src, box, out):
    Image.open(src).convert("RGB").crop(box).save(SHOTS / out)


crop(DOCS / "dashboard-live-vs-shadow.png", (0, 130, 2464, 800), "p_shadow_wide.png")

# ------------------------------------------------------------------ 1 · title
s = slide()
text(s, M, 1.55, CW, 0.35, "ALPACA AI TRADING AGENTS  \u00b7  OPTIONS ALPHA  \u00b7  TEAM AAF11",
     14, bold=True, color=YELLOW)
text(s, M, 1.98, CW, 1.5, "BRIGHTLINE", 92, bold=True)
text(s, M, 3.52, CW, 0.85,
     [("Disciplined AI. Real markets. ", WHITE, False), ("Brighter outcomes.", YELLOW, False)],
     34)
rule(s, 4.5, 3.1)
text(s, M, 4.9, CW, 0.4, "Autonomous options-trading agent  \u00b7  Alpaca paper account",
     19, color=MUTED)
text(s, M, 5.42, CW, 0.4, "Rafan Ahamad Sheik    \u00b7    Prateek D Shriyan", 21)
text(s, M, 5.98, CW, 0.4, URL, 21, bold=True, color=YELLOW, font=MONO)

# ---------------------------------------------------------------- 2 · problem
s = slide()
eyebrow(s, "the problem")
heading(s, "AI can trade. ", "Who controls the trade?")


def _problem(s, x, y, cw, item):
    n, head, body = item
    text(s, x + 0.42, y + 0.38, 0.6, 0.5, n, 30, bold=True, color=YELLOW)
    text(s, x + 0.42, y + 1.02, cw - 0.84, 0.5, head, 25, bold=True)
    text(s, x + 0.42, y + 1.68, cw - 0.84, 1.1, body, 19, color=MUTED, spacing=1.25)


columns(s, [("1", "Hallucination", "It can quote a number\nit was never given."),
            ("2", "Constraint failure", "It can ignore a rule\nit was told to follow."),
            ("3", "Risk failure", "It can size a position\nwith language, not math.")],
        2.35, 2.9, _problem)
footer(s, "Let AI reason about the market. Never let it set the boundaries.")

# --------------------------------------------------------------- 3 · one idea
s = slide()
eyebrow(s, "the one idea")
text(s, M, 2.25, CW, 2.2,
     [("The model may decline a trade it is allowed to take.\n", WHITE, True),
      ("It can never propose one it is not.", YELLOW, True)], 42, spacing=1.25)
rule(s, 4.72, 3.1)
text(s, M, 5.12, CW, 0.9,
     "Not a policy the model is asked to follow \u2014 a property of the code path. The rulebook "
     "is re-derived in Python from the raw signals, never from the model\u2019s summary of them.",
     20, color=MUTED, spacing=1.3)

# ---------------------------------------------------------------- 4 · pipeline
s = slide()
eyebrow(s, "the pipeline")
heading(s, "AI proposes. ", "Python disposes.")
STEPS = (("Signals", "P(Up) \u00b7 IV rank \u00b7 VRP \u00b7 regime \u2014 computed, never asked for"),
         ("Three models", "Identical inputs. Exactly one of them can execute"),
         ("Rulebook", "Re-derived in Python from the raw signals"),
         ("Risk gate", "Quarter-Kelly sizing. Zero AI in the path"),
         ("Execute", "One atomic MLEG order, appended to the audit log"))
for i, (head, body) in enumerate(STEPS):
    y = 2.22 + i * 0.82
    card(s, M, y, 0.62, 0.62, line=YELLOW)
    text(s, M, y + 0.13, 0.62, 0.4, str(i + 1), 22, bold=True, color=YELLOW,
         align=PP_ALIGN.CENTER)
    text(s, M + 0.95, y + 0.06, 2.9, 0.45, head, 24, bold=True)
    text(s, M + 3.9, y + 0.13, CW - 3.9, 0.45, body, 19, color=MUTED)
footer(s, "Every 15 minutes, market hours. Two independent triggers. No human in a session.")

# ---------------------------------------------------------------- 5 · signals
s = slide()
eyebrow(s, "the inputs")
heading(s, "The agent builds ", "its own signals.")


def _signal(s, x, y, cw, item):
    name, body = item
    text(s, x + 0.34, y + 0.42, cw - 0.68, 0.6, name, 27, bold=True, color=YELLOW)
    text(s, x + 0.34, y + 1.18, cw - 0.68, 1.3, body, 18, color=MUTED, spacing=1.25)


columns(s, [("P(Up)", "Calibrated gradient\nboosting on daily bars."),
            ("IV Rank", "Trusted only after 30\nsamples across two days."),
            ("VRP", "Implied volatility\nminus realised."),
            ("Regime", "Volatility state of\nthe underlying.")],
        2.35, 2.9, _signal)
footer(s, "Computed in Python and handed to the model. Never invented by it.")

# ----------------------------------------------------------------- 6 · models
s = slide()
eyebrow(s, "three models")
heading(s, "Same inputs. ", "One execution path.")
MODELS = (("Groq", "LIVE \u00b7 CAN EXECUTE", "Reaches the broker.\nSubject to every check.", True),
          ("Featherless", "SHADOW", "Reasons and is scored.\nNever executes.", False),
          ("Mistral", "SHADOW", "Reasons and is scored.\nNever executes.", False))
cw = (CW - 0.7) / 3
for i, (name, tag, body, live) in enumerate(MODELS):
    x = M + i * (cw + 0.35)
    card(s, x, 2.35, cw, 2.9, line=YELLOW if live else CARD_LINE, width=1.75 if live else 1.0)
    text(s, x + 0.42, 2.72, cw - 0.84, 0.6, name, 30, bold=True)
    text(s, x + 0.42, 3.4, cw - 0.84, 0.4, tag, 14, bold=True,
         color=YELLOW if live else MUTED, font=MONO)
    text(s, x + 0.42, 4.0, cw - 0.84, 1.1, body, 19, color=MUTED, spacing=1.25)
footer(s, "Because the rulebook is a pure function of the same signals, every answer is markable.")

# ----------------------------------------------------------------- 7 · guards
s = slide()
eyebrow(s, "the guards")
heading(s, "Four checks ", "before the rulebook.")
GUARDS = (("Signal validation", "NaN, out of range or inconsistent \u2014 no model is called"),
          ("Faithfulness", "A figure quoted but never supplied kills the trade"),
          ("Cross-model", "Agreement recorded across all three. Never enforced"),
          ("Reconciliation", "A second, independent read of live broker state"))
for i, (head, body) in enumerate(GUARDS):
    y = 2.42 + i * 0.95
    card(s, M, y, 0.62, 0.62, line=YELLOW)
    text(s, M, y + 0.13, 0.62, 0.4, str(i + 1), 22, bold=True, color=YELLOW,
         align=PP_ALIGN.CENTER)
    text(s, M + 0.95, y + 0.02, 3.4, 0.45, head, 24, bold=True)
    text(s, M + 4.4, y + 0.09, CW - 4.4, 0.45, body, 18, color=MUTED)
footer(s, "Five stages can stop a trade. One can start one.")

# ------------------------------------------------------------------- 8 · risk
s = slide()
eyebrow(s, "risk engine")
heading(s, "Risk is ", "not an AI decision.")


# The funnel, drawn rather than screenshotted so it stays legible at any size.
FUNNEL = ("Trade proposal enters", "Quarter-Kelly sizing", "Per-trade and open-risk caps",
          "Drawdown halt \u00b7 loss streak", "Approved   or   rejected")
fx, fw0, fw1 = M, 5.5, 3.2
for i, label in enumerate(FUNNEL):
    y = 2.46 + i * 0.72
    w = fw0 + (fw1 - fw0) * i / (len(FUNNEL) - 1)
    last = i == len(FUNNEL) - 1
    card(s, fx + (fw0 - w) / 2, y, w, 0.6,
         fill=RGBColor(0x10, 0x2A, 0x66), line=YELLOW if last else CARD_LINE,
         width=1.75 if last else 1.0)
    text(s, fx + (fw0 - w) / 2, y, w, 0.6, label, 15, bold=last,
         color=WHITE if last else MUTED, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

LIMITS = (("\u2264 2%", "equity per trade"), ("\u2264 10%", "total open risk"),
          ("\u2264 4", "concurrent positions"), ("1", "position per underlying"))
for i, (num, lab) in enumerate(LIMITS):
    y = 2.5 + i * 0.9
    text(s, 7.1, y, 1.9, 0.7, num, 34, bold=True, color=YELLOW)
    text(s, 9.05, y + 0.16, 3.5, 0.5, lab, 19, color=MUTED)
footer(s, "Every number above is computed in Python. No model can move one of them.")

# --------------------------------------------------------------- 9 · evidence
s = slide()
eyebrow(s, "the evidence")
heading(s, "Measured, ", "not claimed.")


STATS = (("302", "decision cycles"), ("22", "off-book attempts"),
         ("0", "reached the broker"), ("45%", "model disagreement"))
sw = CW / 4
for i, (num, lab) in enumerate(STATS):
    x = M + i * sw
    text(s, x, 2.02, sw, 0.9, num, 50, bold=True, color=YELLOW, align=PP_ALIGN.CENTER)
    text(s, x, 2.98, sw, 0.5, lab, 18, color=MUTED, align=PP_ALIGN.CENTER)
picture(s, "p_equity.png", 3.66, max_bottom=6.15)
footer(s, "Equity is −5.8%. The cause is traced to a units error, and written up in full.")

# --------------------------------------------------------------- 10 · activity
s = slide()
eyebrow(s, "the public record")
heading(s, "Every cycle is ", "logged.")
picture(s, "p_activity.png", 2.3,
        caption="A refusal appears because the gate refused, and it carries the reason it gave.")
footer(s, URL + "  \u2014  read-only, rendered from the audit log.")

# -------------------------------------------------------------- 11 · risk gate
s = slide()
eyebrow(s, "the public record")
heading(s, "Approved, ", "then one atomic order.")
picture(s, "p_riskgate.png", 2.45,
        caption="12 contracts, 1.98% of equity, Kelly-sized \u2014 then a single multi-leg order id.")
footer(s, "Sizing is computed in Python. The model never sees a dollar figure it can change.")

# ----------------------------------------------------------------- 12 · shadow
s = slide()
eyebrow(s, "live vs. shadow")
heading(s, "Three models. ", "Three answers.")
picture(s, "p_shadow_wide.png", 2.4,
        caption="Full reasoning, nothing truncated. Only Groq can act on any of it.")
footer(s, "They disagreed on 45% of cycles. That is the number worth taking away.")

# ------------------------------------------------------------ 13 · what broke
s = slide()
eyebrow(s, "what went wrong")
heading(s, "The P&L is negative. ", "Here is why.")


def _lesson(s, x, y, cw, item):
    head, body = item
    text(s, x + 0.42, y + 0.38, cw - 0.84, 0.5, head, 25, bold=True, color=YELLOW)
    text(s, x + 0.42, y + 1.05, cw - 0.84, 1.5, body, 18, color=MUTED, spacing=1.3)


columns(s, [("A units error",
             "The label is P(3-day return > +0.5%), whose base rate is 0.40\u20130.49. "
             "The rulebook read 0.50 as neutral, so the bullish branch never fired."),
            ("An untrusted number",
             "IV rank looked authoritative before it had earned it. The gate refused 22 "
             "premium sales on it \u2014 which would probably have made money.")],
        2.35, 2.9, _lesson)
footer(s, "Both halves are in the write-up. The gate held; the rulebook it enforced was wrong.")

# ---------------------------------------------------------------- 14 · closing
s = slide()
eyebrow(s, "what it taught us")
text(s, M, 2.15, CW, 1.4, [("Instrumentation ", WHITE, True), ("> assumptions.", YELLOW, True)],
     46, spacing=1.15)
text(s, M, 3.5, CW, 1.2,
     "Every failure above was invisible until it was logged, cross-compared and validated. "
     "The reliability of this system is a property of its observability.",
     21, color=MUTED, spacing=1.35)
rule(s, 5.35, 3.1)
text(s, M, 5.72, CW, 0.4, "Rafan Ahamad Sheik    \u00b7    Prateek D Shriyan    \u00b7    Team AAF11",
     19, color=MUTED)
text(s, M, 6.24, CW, 0.4, URL, 22, bold=True, color=YELLOW, font=MONO)

prs.save("deck14.pptx")
print("slides:", len(prs.slides._sldIdLst))
